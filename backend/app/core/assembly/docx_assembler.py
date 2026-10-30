"""Assemble a complete Word document from classified elements.

Elements are emitted in SOURCE ORDER: text as paragraphs, tables as
native python-docx tables, diagrams via the existing vision → layout →
``docx_shapes`` chain (unmodified). Styling fidelity is intentionally
limited to reasonable defaults — original fonts/sizes are not
reproduced; subscript/superscript runs captured by the extractor
(``metadata["runs"]``) are preserved.
"""

import logging
from io import BytesIO
from typing import Any

import docx
from docx.document import Document
from docx.shared import Emu, Inches
from docx.text.paragraph import Paragraph

from app.core.layout import layout_graph
from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.rich_text import markup_runs
from app.core.vision import get_vision_provider
from app.core.vision.base import VisionProvider
from app.models.document import ClassifiedElement, ElementType
from app.models.text_content import TextContent

logger = logging.getLogger(__name__)

_NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

_DIAGRAM_MAX_W = Emu(int(6.0 * 914400))
_DIAGRAM_MAX_H = Emu(int(7.5 * 914400))


def assemble_docx(
    classified: list[ClassifiedElement],
    vision_provider: VisionProvider | None = None,
) -> Document:
    """Build a complete Document from classified elements, in order."""
    document = docx.Document()
    provider: VisionProvider | None = vision_provider
    last_text: str | None = None

    for item in classified:
        if item.element_type == ElementType.TEXT:
            _add_text(document, item)
            last_text = str(item.element.raw_content)
        elif item.element_type == ElementType.TABLE:
            _add_table(document, item)
        elif item.element_type == ElementType.DIAGRAM:
            if provider is None:
                provider = get_vision_provider()
            _add_vision_element(document, item, provider, context=last_text)
    logger.info("Assembled DOCX with %d elements", len(classified))
    return document


def _emit_runs(paragraph: Paragraph, runs: list[dict[str, str | None]]) -> None:
    """Append run-dicts (``{text, vert_align}``) with subscript formatting.

    Shared by plain text and text-content cells so subscript handling
    lives in one place.
    """
    for run_data in runs:
        run = paragraph.add_run(run_data["text"])
        if run_data.get("vert_align") == "subscript":
            run.font.subscript = True
        elif run_data.get("vert_align") == "superscript":
            run.font.superscript = True


def _add_text(document: Document, item: ClassifiedElement) -> None:
    """Append a paragraph, re-emitting subscript/superscript runs."""
    runs = item.element.metadata.get("runs")
    paragraph = document.add_paragraph()
    if runs:
        _emit_runs(paragraph, runs)
    else:
        paragraph.add_run(str(item.element.raw_content))


def _add_table(document: Document, item: ClassifiedElement) -> None:
    """Append the flattened table text as a native Word table."""
    rows = str(item.element.raw_content).split("\n")
    cells_per_row = [row.split("\t") for row in rows]
    n_cols = int(
        item.element.metadata.get("n_cols")
        or max(len(cells) for cells in cells_per_row)
    )
    table = document.add_table(rows=len(cells_per_row), cols=n_cols)
    table.style = "Table Grid"
    for row_idx, cells in enumerate(cells_per_row):
        for col_idx, text in enumerate(cells[:n_cols]):
            table.cell(row_idx, col_idx).text = text


def _add_vision_element(
    document: Document,
    item: ClassifiedElement,
    provider: VisionProvider,
    context: str | None,
) -> None:
    """Render a vision-analysed image: text content, diagram, or image.

    The single vision call classifies the image; TEXT_CONTENT becomes a
    native 2-column table, a DIAGRAM with edges becomes native shapes,
    and anything else (error, edgeless graph, unusable list structure)
    falls back to embedding the original image.
    """
    image = item.element.raw_content
    assert isinstance(image, bytes)
    result = provider.analyze_diagram(image, context=context)

    if result.text_content is not None and _add_text_content(
        document, result.text_content
    ):
        return

    graph = result.diagram_graph
    if graph is not None and graph.edges:
        if graph.title and not graph.title_below:
            _add_rich_paragraph(document, graph.title)
        layouted = layout_graph(graph)
        add_diagram_to_docx(
            layouted,
            document,
            max_width_emu=_DIAGRAM_MAX_W,
            max_height_emu=_DIAGRAM_MAX_H,
            relative_from="paragraph",
        )
        _flow_text_around(document)
        if graph.title and graph.title_below:
            _add_rich_paragraph(document, graph.title)
        return

    logger.warning(
        "No usable diagram/text for element %s (%s); embedding original image",
        item.element.id,
        result.error or "degenerate content",
    )
    document.add_picture(BytesIO(image), width=Inches(4.5))


def _add_rich_paragraph(document: Document, text: str) -> None:
    """A paragraph honouring _{...}/^{...} markup as sub/superscript runs."""
    paragraph = document.add_paragraph()
    for chunk, vert in markup_runs(text):
        run = paragraph.add_run(chunk)
        # Both flags share one w:vertAlign element: set only the one needed.
        if vert == "sub":
            run.font.subscript = True
        elif vert == "sup":
            run.font.superscript = True


def _remove_table_borders(table: Any) -> None:
    """Set all tblBorders edges to "nil" so no border/gridline prints.

    Scoped to TEXT_CONTENT tables — a matching question has no visible
    grid in the source scan (students draw connecting lines by hand),
    unlike the real data tables which keep their normal borders.
    """
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(f"{{{_NS_W}}}tblBorders", {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(f"{{{_NS_W}}}{edge}", {f"{{{_NS_W}}}val": "nil"})
        borders.append(el)
    # CT_TblPr requires schema order …tblW, jc, tblCellSpacing, tblInd,
    # tblBorders, shd, tblLayout, tblCellMar, tblLook…; insert before the
    # first element that must follow tblBorders (tblLook, in practice).
    successor = tbl_pr.find(f"{{{_NS_W}}}tblLook")
    if successor is not None:
        successor.addprevious(borders)
    else:
        tbl_pr.append(borders)


def _add_text_content(document: Document, text_content: TextContent) -> bool:
    """Render a two-list matching question as a native, borderless table.

    A heading paragraph (if present) is emitted directly above the
    table. Left column = first list, right column = second list, one
    row per line with the shorter column padded by empty cells. Returns
    False (rendering nothing) when the content is not a usable two-list
    structure, so the caller can fall back to the original image.
    """
    columns = text_content.usable_columns()
    if columns is None:
        return False
    if text_content.heading:
        document.add_paragraph().add_run(text_content.heading)
    left, right = columns
    n_rows = max(len(left.lines), len(right.lines))
    table = document.add_table(rows=n_rows, cols=2)
    table.style = "Table Grid"
    _remove_table_borders(table)
    for col_idx, column in enumerate((left, right)):
        for row_idx, line in enumerate(column.lines):
            cell = table.cell(row_idx, col_idx)
            # A fresh cell has one empty paragraph; fill its runs.
            _emit_runs(cell.paragraphs[0], line.to_run_dicts())
    return True


def _flow_text_around(document: Document) -> None:
    """Switch the just-inserted diagram anchor to top-and-bottom wrapping.

    ``docx_shapes`` emits ``wp:wrapNone`` (fine for a standalone
    diagram); in an assembled document the following paragraphs must
    flow BELOW the canvas, not underneath it. Post-processing here keeps
    ``docx_shapes.py`` untouched.
    """
    anchors = document.element.body.findall(f".//{{{_NS_WP}}}anchor")
    anchor = anchors[-1]
    wrap_none = anchor.find(f"{{{_NS_WP}}}wrapNone")
    if wrap_none is not None:
        wrap = anchor.makeelement(f"{{{_NS_WP}}}wrapTopAndBottom", {})
        anchor.replace(wrap_none, wrap)
