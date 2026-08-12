"""One-shot Phase 3 validation: extract → classify → vision → layout →
native PPTX shapes (+ debug SVG).

Runs the full first-to-third-stage pipeline on the real flowchart sample
and writes samples/output/flowchart_diagram.pptx (client-style artifact,
native editable shapes) and samples/output/flowchart_debug.svg (internal
QA preview). Requires GEMINI_API_KEY in backend/.env.

Usage (from backend/):
    venv/Scripts/python scripts/validate_pipeline.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx  # noqa: E402
from docx.shared import Emu  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from app.config import get_settings  # noqa: E402
from app.core.classifier import classify_elements  # noqa: E402
from app.core.ingestion import get_extractor  # noqa: E402
from app.core.layout import layout_graph  # noqa: E402
from app.core.render.docx_shapes import add_diagram_to_docx  # noqa: E402
from app.core.render.pptx_shapes import add_diagram_to_slide  # noqa: E402
from app.core.render.svg_debug import render_svg  # noqa: E402
from app.core.vision import get_vision_provider  # noqa: E402
from app.logging_config import configure_logging  # noqa: E402
from app.models.document import ElementKind, ElementType  # noqa: E402

SAMPLES = Path(__file__).resolve().parents[2] / "samples"
OUTPUT = SAMPLES / "output"
FLOWCHART = (SAMPLES / "flowchart_a3.png").read_bytes()


def main() -> int:
    """Run the extraction → classification → vision → render chain once."""
    configure_logging(get_settings().log_level)
    OUTPUT.mkdir(exist_ok=True)

    source = SAMPLES / "informatika_seh_40_41.docx"
    classified = classify_elements(get_extractor(source).extract(source))

    flowchart_idx = next(
        i for i, c in enumerate(classified) if c.element.raw_content == FLOWCHART
    )
    item = classified[flowchart_idx]
    assert item.element_type == ElementType.DIAGRAM
    context = next(
        (
            str(c.element.raw_content)
            for c in reversed(classified[:flowchart_idx])
            if c.element.element_kind == ElementKind.TEXT_BLOCK
        ),
        None,
    )
    print(f"Context passed to the model: {context!r}\n")

    image = item.element.raw_content
    assert isinstance(image, bytes)
    result = get_vision_provider().analyze_diagram(image, context=context)
    if result.error or result.diagram_graph is None:
        print(f"ERROR: {result.error}")
        print(f"Raw model output:\n{result.raw_model_output}")
        return 1

    graph = result.diagram_graph
    print("=== DIAGRAM GRAPH " + "=" * 42)
    for node in graph.nodes:
        print(f"  {node.id}: [{node.shape_type.value}] {node.text!r}")
    for edge in graph.edges:
        label = f" ({edge.label})" if edge.label else ""
        print(f"  {edge.source_id} -> {edge.target_id}{label}")
    print("=" * 60)

    layouted = layout_graph(graph)
    back_edges = [e for e in layouted.edges if e.is_back_edge]
    print(
        f"Layout: {len(layouted.nodes)} nodes, {len(layouted.edges)} edges, "
        f"{len(back_edges)} back-edge(s): "
        + ", ".join(f"{e.source_id}->{e.target_id}" for e in back_edges)
    )

    svg_path = OUTPUT / "flowchart_debug.svg"
    svg_path.write_text(render_svg(layouted), encoding="utf-8")
    print(f"Saved debug SVG (internal QA only): {svg_path}")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    add_diagram_to_slide(
        layouted,
        slide,
        offset_x_emu=Inches(0.5),
        offset_y_emu=Inches(0.2),
        max_width_emu=Inches(12),
        max_height_emu=Inches(7),
    )
    pptx_path = OUTPUT / "flowchart_diagram.pptx"
    prs.save(str(pptx_path))
    print(f"Saved native-shape PPTX (open by hand to verify!): {pptx_path}")

    document = docx.Document()
    add_diagram_to_docx(
        layouted,
        document,
        offset_x_emu=Emu(0),
        offset_y_emu=Emu(0),
        max_width_emu=Emu(int(6.5 * 914400)),  # letter width minus 1" margins
        max_height_emu=Emu(int(8.5 * 914400)),
        relative_from="paragraph",
    )
    docx_path = OUTPUT / "flowchart_diagram.docx"
    document.save(str(docx_path))
    print(f"Saved native-shape DOCX (open by hand in Word to verify!): {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
