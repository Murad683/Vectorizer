"""Phase 23: real Gemini verification — 1 API call on Blok Sxem1.jpeg."""

import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from lxml import etree
from pptx import Presentation
from pptx.shapes.connector import Connector
from pptx.util import Inches
import docx
from docx.shared import Emu

from app.core.layout import layout_graph
from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.core.vision import get_vision_provider

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"

SAMPLE = Path(__file__).resolve().parents[2] / "samples" / "Blok Sxem1.jpeg"
OUT = Path(__file__).resolve().parents[2] / "samples" / "output"


def main():
    OUT.mkdir(exist_ok=True)
    print(f"Reading: {SAMPLE}")
    image_bytes = SAMPLE.read_bytes()

    print("Calling Gemini (1 real API call)...")
    provider = get_vision_provider()
    result = provider.analyze_diagram(image_bytes)

    if result.error:
        print(f"ERROR: {result.error}")
        return False
    if result.diagram_graph is None:
        print("ERROR: No diagram_graph returned")
        return False

    graph = result.diagram_graph
    print(f"Diagram: {len(graph.nodes)} nodes, {len(graph.edges)} edges")
    if graph.title:
        print(f"  Title: {graph.title}")
    for n in graph.nodes:
        print(f"  {n.id}: {n.shape_type.value} \"{n.text}\"")
    for e in graph.edges:
        lbl = f" [{e.label}]" if e.label else ""
        print(f"  {e.source_id}->{e.target_id}{lbl}")

    print("\nRunning layout (Graphviz dot ortho)...")
    layouted = layout_graph(graph)

    # Generate PPTX
    pptx_path = OUT / "phase23_verify.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(
        layouted, slide,
        offset_x_emu=Inches(0.5), offset_y_emu=Inches(0.2),
        max_width_emu=Inches(9), max_height_emu=Inches(7),
    )
    prs.save(str(pptx_path))
    print(f"PPTX saved: {pptx_path}")

    # Generate DOCX
    docx_path = OUT / "phase23_verify.docx"
    document = docx.Document()
    add_diagram_to_docx(
        layouted, document,
        max_width_emu=Emu(Inches(6)), max_height_emu=Emu(Inches(8)),
    )
    document.save(str(docx_path))
    print(f"DOCX saved: {docx_path}")

    # Verify both outputs
    ok = True

    print("\n=== PPTX Verification ===")
    prs2 = Presentation(str(pptx_path))
    ns = {"a": NS_A}
    conn_count = 0
    for sl in prs2.slides:
        for shape in sl.shapes:
            if not isinstance(shape, Connector):
                continue
            conn_count += 1
            prst_el = shape._element.find(".//a:prstGeom", ns)
            if prst_el is None:
                continue
            prst = prst_el.get("prst", "")
            xfrm = shape._element.spPr.find(f"{{{NS_A}}}xfrm")
            ext = xfrm.find(f"{{{NS_A}}}ext") if xfrm is not None else None
            cx = int(ext.get("cx", "0")) if ext is not None else 0
            cy = int(ext.get("cy", "0")) if ext is not None else 0
            st = shape._element.find(".//a:stCxn", ns) is not None
            en = shape._element.find(".//a:endCxn", ns) is not None
            diag = prst == "straightConnector1" and abs(cx) > 1 and abs(cy) > 1
            if diag:
                ok = False
                print(f"  DIAGONAL! prst={prst} cx={cx} cy={cy}")
            elif not (st and en):
                ok = False
                print(f"  UNGLUED! prst={prst}")
            else:
                print(f"  OK prst={prst} cx={cx} cy={cy}")
    print(f"  Total connectors: {conn_count}")

    print("\n=== DOCX Verification ===")
    with zipfile.ZipFile(docx_path) as z:
        doc_xml = etree.fromstring(z.read("word/document.xml"))
    conn_count = 0
    for wsp in doc_xml.iter(f"{{{NS_WPS}}}wsp"):
        if wsp.find(f"{{{NS_WPS}}}cNvCnPr") is None:
            continue
        conn_count += 1
        sp_pr = wsp.find(f"{{{NS_WPS}}}spPr")
        if sp_pr is None:
            continue
        prst_geom = sp_pr.find(f"{{{NS_A}}}prstGeom")
        if prst_geom is None:
            continue
        prst = prst_geom.get("prst", "")
        xfrm = sp_pr.find(f"{{{NS_A}}}xfrm")
        ext = xfrm.find(f"{{{NS_A}}}ext") if xfrm is not None else None
        cx = int(ext.get("cx", "0")) if ext is not None else 0
        cy = int(ext.get("cy", "0")) if ext is not None else 0
        cn_pr = wsp.find(f"{{{NS_WPS}}}cNvCnPr")
        st = cn_pr.find(f"{{{NS_A}}}stCxn") is not None if cn_pr is not None else False
        en = cn_pr.find(f"{{{NS_A}}}endCxn") is not None if cn_pr is not None else False
        diag = prst == "straightConnector1" and abs(cx) > 1 and abs(cy) > 1
        if diag:
            ok = False
            print(f"  DIAGONAL! prst={prst} cx={cx} cy={cy}")
        elif not (st and en):
            ok = False
            print(f"  UNGLUED! prst={prst}")
        else:
            print(f"  OK prst={prst} cx={cx} cy={cy}")
    print(f"  Total connectors: {conn_count}")

    print(f"\n{'PASS' if ok else 'FAIL'} — all connectors checked")
    return ok


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
