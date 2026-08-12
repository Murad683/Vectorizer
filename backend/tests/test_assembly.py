"""Assembly tests: full-document DOCX/PPTX from the real sample (vision mocked)."""

import zipfile
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from lxml import etree
from pptx.util import Emu

from app.core.assembly import assemble_docx, assemble_pptx
from app.core.classifier import classify_elements
from app.core.ingestion.docx_extractor import DocxExtractor
from app.models.diagram import DiagramGraph
from app.models.document import ClassifiedElement, VisionResult

_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": ("http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"),
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
}


@pytest.fixture()
def classified(sample_docx: Path) -> list[ClassifiedElement]:
    return classify_elements(DocxExtractor().extract(sample_docx))


@pytest.fixture()
def mock_vision(flowchart_graph: DiagramGraph, flowchart_png_bytes: bytes) -> MagicMock:
    """Vision that solves the real flowchart and fails on other images."""

    def analyze(image: bytes, context: str | None = None) -> VisionResult:
        if image == flowchart_png_bytes:
            return VisionResult(diagram_graph=flowchart_graph)
        return VisionResult(error="not a flowchart (mock)")

    provider = MagicMock()
    provider.analyze_diagram.side_effect = analyze
    return provider


def _parse_package(path: Path) -> None:
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.endswith((".xml", ".rels")):
                etree.fromstring(zf.read(name))


class TestDocxAssembly:
    def test_full_document_contains_all_element_kinds_in_order(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
        tmp_path: Path,
    ) -> None:
        document = assemble_docx(classified, vision_provider=mock_vision)
        path = tmp_path / "full.docx"
        document.save(str(path))
        _parse_package(path)  # every XML part parses

        body = document.element.body
        # Body children preserve source order: the diagram canvas sits
        # right after question 80's paragraph, before the answers.
        children = list(body)
        texts_or_drawings = []
        for child in children:
            if child.tag == f"{{{_NS['w']}}}p":
                if child.findall(f".//{{{_NS['wpc']}}}wpc"):
                    texts_or_drawings.append("<canvas>")
                else:
                    text = "".join(
                        t.text or "" for t in child.findall(f".//{{{_NS['w']}}}t")
                    )
                    texts_or_drawings.append(text)
            elif child.tag == f"{{{_NS['w']}}}tbl":
                texts_or_drawings.append("<table>")
        q80 = next(i for i, t in enumerate(texts_or_drawings) if t.startswith("80."))
        assert texts_or_drawings[q80 + 1] == "<canvas>"
        assert texts_or_drawings[q80 + 2].startswith("A) 320")
        # All three source tables present as native tables.
        assert texts_or_drawings.count("<table>") == 3

    def test_tables_are_native_with_correct_shape(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
    ) -> None:
        document = assemble_docx(classified, vision_provider=mock_vision)
        tables = document.tables
        assert len(tables) == 3
        marka = tables[0]
        assert marka.cell(0, 1).text == "Marka"
        assert marka.cell(1, 1).text == "Mercedes-Benz"

    def test_subscript_runs_preserved(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
    ) -> None:
        document = assemble_docx(classified, vision_provider=mock_vision)
        # Question 81's "41x + 36x = 77x" — x's are subscript runs.
        para = next(p for p in document.paragraphs if p.text.startswith("81."))
        sub_runs = [r.text for r in para.runs if r.font.subscript]
        assert sub_runs == ["x", "x", "x"]

    def test_failed_vision_falls_back_to_picture(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
    ) -> None:
        document = assemble_docx(classified, vision_provider=mock_vision)
        # 2 DIAGRAM elements: flowchart → canvas; table_87 image → mock
        # error → embedded picture fallback.
        canvases = document.element.body.findall(f".//{{{_NS['wpc']}}}wpc")
        assert len(canvases) == 1
        drawings_with_pics = document.element.body.findall(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"
        )
        assert len(drawings_with_pics) == 1

    def test_diagram_anchor_wraps_top_and_bottom(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
    ) -> None:
        document = assemble_docx(classified, vision_provider=mock_vision)
        anchor = document.element.body.find(f".//{{{_NS['wp']}}}anchor")
        assert anchor.find(f"{{{_NS['wp']}}}wrapNone") is None
        assert anchor.find(f"{{{_NS['wp']}}}wrapTopAndBottom") is not None


class TestPptxAssembly:
    def test_full_presentation_contains_everything(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
        tmp_path: Path,
    ) -> None:
        prs = assemble_pptx(classified, vision_provider=mock_vision)
        path = tmp_path / "full.pptx"
        prs.save(str(path))
        _parse_package(path)  # every XML part parses

        slides = list(prs.slides)
        assert len(slides) > 1  # flow-fill produced multiple slides
        all_text = [
            shape.text_frame.text
            for slide in slides
            for shape in slide.shapes
            if shape.has_text_frame
        ]
        assert any(t.startswith("77.") for t in all_text)  # first question
        assert any("Başlanğıc" in t for t in all_text)  # diagram shapes
        tables = [
            shape for slide in slides for shape in slide.shapes if shape.has_table
        ]
        assert len(tables) == 3
        pictures = [
            shape
            for slide in slides
            for shape in slide.shapes
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type)
        ]
        assert len(pictures) == 1  # table_87 fallback image

    def test_diagram_gets_its_own_slide_in_sequence(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
    ) -> None:
        prs = assemble_pptx(classified, vision_provider=mock_vision)
        diagram_slides = [
            slide
            for slide in prs.slides
            if any(
                shape.has_text_frame and shape.text_frame.text == "Başlanğıc"
                for shape in slide.shapes
            )
        ]
        assert len(diagram_slides) == 1
        slide = diagram_slides[0]
        # Nothing but the diagram's shapes/connectors/labels on that slide.
        assert all(not shape.has_table for shape in slide.shapes)

    def test_slide_dimensions(
        self,
        classified: list[ClassifiedElement],
        mock_vision: MagicMock,
    ) -> None:
        prs = assemble_pptx(classified, vision_provider=mock_vision)
        assert prs.slide_width == Emu(int(10 * 914400))
        assert prs.slide_height == Emu(int(7.5 * 914400))


def test_subscript_runs_captured_by_extractor(sample_docx: Path) -> None:
    elements = DocxExtractor().extract(sample_docx)
    q81 = next(e for e in elements if str(e.raw_content).startswith("81. 41x"))
    runs = q81.metadata.get("runs")
    assert runs is not None
    subs = [r["text"] for r in runs if r.get("vert_align") == "subscript"]
    assert subs == ["x", "x", "x"]
