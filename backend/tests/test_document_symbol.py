"""The flowchart Document symbol (print block with a wavy bottom), multi-line
node text, captions below the diagram and the prompt rules behind them.
All offline."""

import io
import zipfile

import docx
import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.core.diagram_detect import apply_detection, detect
from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.core.shape_geometry import document_bottom
from app.core.vision.gemini_provider import _PROMPT
from app.core.vision.postprocess import topology_problems
from app.models.diagram import DiagramEdge, DiagramGraph, DiagramNode, ShapeType
from tests.diagram_fixtures import FIXTURES, SAMPLES, run_fixture, simulate_marks

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.mark.parametrize(
    "file,documents",
    [
        ("quadratic_equation.png", [6, 7, 8]),
        ("sum_loop.jpg", [8]),
        ("fahrenheit_celsius.png", []),  # its PRINT block is a parallelogram
        ("sp_open_exit.png", []),
        ("Blok Sxem1.jpeg", []),
        ("flowchart_a3.png", []),
    ],
)
def test_detector_recognises_wavy_bottom(file: str, documents: list[int]) -> None:
    det = detect((SAMPLES / file).read_bytes())
    assert det is not None
    found = [i for i, s in enumerate(det.shapes, start=1) if s.kind == "document"]
    assert found == documents


def test_measured_outline_overrides_model_type() -> None:
    """A model that calls the print block a parallelogram is corrected; a
    type the outline says nothing about is left alone."""
    file, make = FIXTURES["sum_loop"]
    image = (SAMPLES / file).read_bytes()
    det = detect(image)
    assert det is not None
    graph = make()
    printer = next(n for n in graph.nodes if n.id == "n8")
    printer.shape_type = ShapeType.INPUT_OUTPUT
    simulate_marks(graph, det)
    apply_detection(graph, det)
    assert printer.shape_type == ShapeType.DOCUMENT
    reader = next(n for n in graph.nodes if n.id == "n2")
    assert reader.shape_type == ShapeType.INPUT_OUTPUT


def test_arrow_leaves_from_the_wave_not_the_empty_corner() -> None:
    _, lg = run_fixture("sum_loop")
    node = lg.node_by_id("n8")
    edge = next(e for e in lg.edges if e.source_id == "n8")
    assert edge.edge_path is not None
    x, y = edge.edge_path[0]
    fx = (x - node.x) / node.width
    assert y == pytest.approx(node.y + node.height * document_bottom(fx), abs=0.5)
    assert y < node.y + node.height - 1


def test_document_bottom_matches_preset_endpoints() -> None:
    assert document_bottom(1.0) == pytest.approx(17322 / 21600, abs=1e-3)
    assert document_bottom(0.0) == pytest.approx(20172 / 21600, abs=1e-3)
    assert document_bottom(0.5) == pytest.approx(0.933, abs=2e-3)


def _docx_xml(name: str) -> etree._Element:
    _, lg = run_fixture(name)
    document = docx.Document()
    add_diagram_to_docx(lg, document, max_width_emu=Inches(6), max_height_emu=Inches(7))
    buf = io.BytesIO()
    document.save(buf)
    with zipfile.ZipFile(buf) as z:
        return etree.fromstring(z.read("word/document.xml"))


def test_docx_uses_word_document_preset_and_keeps_line_breaks() -> None:
    root = _docx_xml("quadratic_equation")
    presets = [g.get("prst") for g in root.iter(f"{_A}prstGeom")]
    assert presets.count("flowChartDocument") == 3
    # "x1 və x2-nin qiymətlərinin\nçap edilməsi" -> two paragraphs.
    texts = []
    for box in root.iter(f"{_W}txbxContent"):
        paras = [
            "".join(t.text or "" for t in p.iter(f"{_W}t")) for p in box.iter(f"{_W}p")
        ]
        texts.append(paras)
    assert ["x1 və x2-nin qiymətlərinin", "çap edilməsi"] in texts
    assert [""] in texts  # the empty block stays empty


def test_pptx_uses_flowchart_document_autoshape() -> None:
    _, lg = run_fixture("sum_loop")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(lg, slide, max_width_emu=Inches(9), max_height_emu=Inches(7))
    kinds = [s.auto_shape_type for s in slide.shapes if s.shape_type == 1]  # AUTO_SHAPE
    assert kinds.count(MSO_SHAPE.FLOWCHART_DOCUMENT) == 1
    reader = next(s for s in slide.shapes if s.name.startswith("Parallelogram"))
    lines = [p.text for p in reader.text_frame.paragraphs]
    assert lines == ["X1, X2, ... X20 ədədlərinin", "daxil edilməsi"]


def test_docx_caption_below_follows_the_diagram() -> None:
    from app.core.assembly.docx_assembler import assemble_docx
    from app.models.document import (
        ClassifiedElement,
        ElementKind,
        ElementType,
        ExtractedElement,
    )
    from scripts.replay_diagram import _FixedProvider

    graph, _ = run_fixture("sum_loop")
    element = ClassifiedElement(
        element=ExtractedElement(
            source_page=1,
            element_kind=ElementKind.IMAGE,
            raw_content=(SAMPLES / "sum_loop.jpg").read_bytes(),
        ),
        element_type=ElementType.DIAGRAM,
        rule="test",
    )
    document = assemble_docx([element], vision_provider=_FixedProvider(graph))
    body = list(document.element.body.iterchildren(f"{_W}p"))
    has_drawing = [p.find(f".//{_W}drawing") is not None for p in body]
    texts = ["".join(t.text or "" for t in p.iter(f"{_W}t")) for p in body]
    drawing_at = has_drawing.index(True)
    caption_at = next(i for i, t in enumerate(texts) if t.startswith("Sxem 3."))
    assert caption_at > drawing_at


def test_document_needs_exactly_one_outgoing_arrow() -> None:
    graph = DiagramGraph(
        nodes=[
            DiagramNode(id="n1", shape_type=ShapeType.TERMINATOR, text="Başlanğıc"),
            DiagramNode(id="n2", shape_type=ShapeType.DOCUMENT, text="Çap"),
        ],
        edges=[DiagramEdge(source_id="n1", target_id="n2")],
    )
    assert any("0 outgoing" in p for p in topology_problems(graph))


def test_prompt_states_the_new_rules() -> None:
    assert '"document"' in _PROMPT and "wave" in _PROMPT
    assert '"title_below"' in _PROMPT
    assert 'join them with "\\n"' in _PROMPT
    assert "Never invent" in _PROMPT
    assert "ON ANOTHER ARROW'S LINE" in _PROMPT
