"""Frameless nodes (formulas drawn without an outline) keep exact geometry.

samples/quadratic_equation.png: ``x = -b / 2a`` and the two-root formula
have no frame, so the detector cannot mark them. Before, that dropped the
whole diagram into approximate mode (no original lines as guides), which
produced the crossing seen in Word.
"""

import io
import zipfile

import docx
import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app.core.diagram_detect import apply_detection, detect
from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.core.vision.gemini_provider import _PROMPT
from tests.diagram_fixtures import (
    SAMPLES,
    quadratic_equation,
    run_fixture,
    simulate_marks,
)

_WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_W, _H = 505, 614
# Hand-measured ink extents of the two formulas (x0, y0, x1, y1).
_TRUTH = {"n7": (130, 320, 292, 345), "n9": (117, 420, 301, 467)}


def _px(box: list[int]) -> tuple[float, float, float, float]:
    return (
        box[1] / 1000 * _W,
        box[0] / 1000 * _H,
        box[3] / 1000 * _W,
        box[2] / 1000 * _H,
    )


def _iou(a: tuple[float, ...], b: tuple[float, ...]) -> float:
    ix = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    return inter / (
        (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    )


def _apply(shift: tuple[int, int] = (0, 0)) -> dict[str, tuple[float, ...]]:
    image = (SAMPLES / "quadratic_equation.png").read_bytes()
    det = detect(image)
    assert det is not None
    graph = quadratic_equation()
    simulate_marks(graph, det)
    for node in graph.nodes:
        if not node.framed and node.box_2d:  # model boxes drift: shift them
            dy, dx = round(shift[1] / _H * 1000), round(shift[0] / _W * 1000)
            y0, x0, y1, x1 = node.box_2d
            node.box_2d = [y0 + dy, x0 + dx, y1 + dy, x1 + dx]
    apply_detection(graph, det)
    assert graph.boxes_detected and graph.guide_segments
    return {n.id: _px(n.box_2d) for n in graph.nodes if not n.framed and n.box_2d}


@pytest.mark.parametrize("shift", [(0, 0), (8, 3), (-6, -4), (0, -12)])
def test_frameless_boxes_are_measured_on_the_ink(shift: tuple[int, int]) -> None:
    measured = _apply(shift)
    for nid, truth in _TRUTH.items():
        assert _iou(measured[nid], truth) > 0.8, (nid, measured[nid])


def test_margin_numbers_are_not_part_of_the_formula() -> None:
    measured = _apply()
    assert measured["n7"][0] > 120  # the row number "7" sits at x ~106-112
    assert measured["n9"][0] > 110  # "9" at x ~98-104


def test_unmatched_framed_symbol_still_means_approximate_mode() -> None:
    image = (SAMPLES / "quadratic_equation.png").read_bytes()
    det = detect(image)
    assert det is not None
    graph = quadratic_equation()
    simulate_marks(graph, det)
    graph.nodes[0].mark = None  # the model forgot a framed symbol's mark
    apply_detection(graph, det)
    assert not graph.boxes_detected


def test_frameless_node_is_text_only_in_docx_and_pptx() -> None:
    _, lg = run_fixture("quadratic_equation")
    document = docx.Document()
    add_diagram_to_docx(lg, document, max_width_emu=Inches(6), max_height_emu=Inches(7))
    buf = io.BytesIO()
    document.save(buf)
    with zipfile.ZipFile(buf) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    for wsp in root.iter(f"{_WPS}wsp"):
        name = wsp.find(f"{_WPS}cNvPr").get("name")
        if name in ("node-n7", "node-n9"):
            sp = wsp.find(f"{_WPS}spPr")
            assert sp.find(f"{_A}noFill") is not None
            assert sp.find(f"{_A}ln/{_A}noFill") is not None
        elif name == "node-n8":
            assert wsp.find(f"{_WPS}spPr/{_A}ln/{_A}solidFill") is not None

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(lg, slide, max_width_emu=Inches(9), max_height_emu=Inches(7))
    formula = next(
        s
        for s in slide.shapes
        if s.has_text_frame and s.text_frame.text == "x = (-b)/(2a)"
    )
    assert formula.line.fill.type == 5  # MSO_FILL.BACKGROUND: no outline


def test_prompt_asks_for_frameless_nodes() -> None:
    assert '"framed": false' in _PROMPT


def test_text_in_an_empty_block_is_dropped() -> None:
    """Live answer: the model filled the empty block 3 with D = b^2 - 4ac."""
    image = (SAMPLES / "quadratic_equation.png").read_bytes()
    det = detect(image)
    assert det is not None
    assert [s.empty for s in det.shapes].count(True) == 1
    graph = quadratic_equation()
    block = next(n for n in graph.nodes if n.id == "n3")
    block.text = "D = b^{2} - 4ac"
    simulate_marks(graph, det)
    apply_detection(graph, det)
    assert block.text == "" and block.latex is None
    assert next(n for n in graph.nodes if n.id == "n4").text == "D > 0"
