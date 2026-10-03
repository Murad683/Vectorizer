"""API-level tests for the /jobs endpoints (with auth)."""

import io
import zipfile
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path
from unittest.mock import MagicMock, patch
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.core.auth.models import Role
from app.core.auth.security import hash_password
from app.core.auth.store import RefreshTokenStore, UserStore
from app.core.queue.models import JobStatus
from app.core.queue.store import JobStore
from app.main import create_app


@asynccontextmanager
async def _noop_lifespan(app: FastAPI) -> AsyncIterator[None]:
    yield


@pytest.fixture()
def _stores(
    job_store: JobStore, user_store: UserStore, refresh_token_store: RefreshTokenStore
):
    return job_store, user_store, refresh_token_store


@pytest.fixture()
def client(_stores, tmp_path: Path) -> TestClient:
    job_store, user_store, refresh_token_store = _stores
    admin_id = uuid4()
    user_store.create_user(
        user_id=admin_id,
        email="admin@test.com",
        hashed_password=hash_password("adminpass"),
        role=Role.ADMIN,
    )
    app = create_app()
    app.router.lifespan_context = _noop_lifespan
    app.state.job_store = job_store
    app.state.user_store = user_store
    app.state.refresh_token_store = refresh_token_store
    with TestClient(app, raise_server_exceptions=False) as c:
        yield c


def _login(
    client: TestClient, email: str = "admin@test.com", password: str = "adminpass"
) -> str:
    resp = client.post("/auth/login", json={"email": email, "password": password})
    assert resp.status_code == 200, resp.text
    return resp.json()["access_token"]


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture()
def sample_bytes(sample_docx: Path) -> bytes:
    return sample_docx.read_bytes()


class TestUploadValidation:
    def test_unsupported_format_rejected(self, client: TestClient) -> None:
        token = _login(client)
        resp = client.post(
            "/jobs",
            files={"file": ("test.txt", b"hello", "text/plain")},
            headers=_auth(token),
        )
        assert resp.status_code == 400
        assert resp.json()["error_code"] == "unsupported_format"

    def test_missing_filename_rejected(self, client: TestClient) -> None:
        token = _login(client)
        resp = client.post(
            "/jobs",
            files={"file": ("", b"hello", "application/octet-stream")},
            headers=_auth(token),
        )
        assert resp.status_code in (400, 422)

    def test_oversized_file_rejected(self, client: TestClient) -> None:
        token = _login(client)
        with patch("app.api.routes.jobs.get_settings") as mock_settings:
            settings = MagicMock()
            settings.max_upload_size_mb = 0
            mock_settings.return_value = settings
            resp = client.post(
                "/jobs",
                files={"file": ("test.docx", b"x" * 100, "application/octet-stream")},
                headers=_auth(token),
            )
        assert resp.status_code == 413
        assert resp.json()["error_code"] == "file_too_large"

    def test_unauthenticated_upload_401(self, client: TestClient) -> None:
        resp = client.post(
            "/jobs",
            files={"file": ("test.docx", b"hello", "application/octet-stream")},
        )
        assert resp.status_code == 401

    def test_png_upload_accepted(
        self, client: TestClient, flowchart_png_bytes: bytes
    ) -> None:
        token = _login(client)
        resp = client.post(
            "/jobs",
            files={"file": ("photo.png", flowchart_png_bytes, "image/png")},
            headers=_auth(token),
        )
        assert resp.status_code == 201
        assert resp.json()["status"] == "pending"

    def test_jpeg_upload_accepted(
        self, client: TestClient, flowchart_png_bytes: bytes
    ) -> None:
        token = _login(client)
        resp = client.post(
            "/jobs",
            # Real bytes needn't be true JPEG-encoded for this endpoint —
            # it's an extension-based check, matching the existing PDF/DOCX
            # behavior (no content sniffing at upload time; see the
            # ImageExtractor tests for the actual-content validation,
            # which runs during background processing instead).
            files={"file": ("photo.jpg", flowchart_png_bytes, "image/jpeg")},
            headers=_auth(token),
        )
        assert resp.status_code == 201
        assert resp.json()["status"] == "pending"

    def test_oversized_image_rejected(
        self, client: TestClient, flowchart_png_bytes: bytes
    ) -> None:
        token = _login(client)
        with patch("app.api.routes.jobs.get_settings") as mock_settings:
            settings = MagicMock()
            settings.max_upload_size_mb = 0
            mock_settings.return_value = settings
            resp = client.post(
                "/jobs",
                files={"file": ("photo.png", flowchart_png_bytes, "image/png")},
                headers=_auth(token),
            )
        assert resp.status_code == 413
        assert resp.json()["error_code"] == "file_too_large"


class TestJobLifecycle:
    def test_create_returns_pending(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        token = _login(client)
        resp = client.post(
            "/jobs",
            files={"file": ("test.docx", sample_bytes, "application/octet-stream")},
            headers=_auth(token),
        )
        assert resp.status_code == 201
        body = resp.json()
        assert body["status"] == "pending"
        assert "job_id" in body

    def test_get_status_pending(self, client: TestClient, sample_bytes: bytes) -> None:
        token = _login(client)
        create_resp = client.post(
            "/jobs",
            files={"file": ("test.docx", sample_bytes, "application/octet-stream")},
            headers=_auth(token),
        )
        job_id = create_resp.json()["job_id"]
        resp = client.get(f"/jobs/{job_id}", headers=_auth(token))
        assert resp.status_code == 200
        assert resp.json()["status"] == "pending"

    def test_get_nonexistent_job_404(self, client: TestClient) -> None:
        token = _login(client)
        resp = client.get(
            "/jobs/00000000-0000-0000-0000-000000000000",
            headers=_auth(token),
        )
        assert resp.status_code == 404

    def test_download_before_completion_409(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        token = _login(client)
        create_resp = client.post(
            "/jobs",
            files={"file": ("test.docx", sample_bytes, "application/octet-stream")},
            headers=_auth(token),
        )
        job_id = create_resp.json()["job_id"]
        resp = client.get(f"/jobs/{job_id}/download?format=docx", headers=_auth(token))
        assert resp.status_code == 409

    def test_download_invalid_format_400(
        self, client: TestClient, sample_bytes: bytes, tmp_path: Path
    ) -> None:
        token = _login(client)
        create_resp = client.post(
            "/jobs",
            files={"file": ("test.docx", sample_bytes, "application/octet-stream")},
            headers=_auth(token),
        )
        job_id = create_resp.json()["job_id"]
        from uuid import UUID

        store = client.app.state.job_store  # type: ignore[union-attr]
        store.update_job(UUID(job_id), status=JobStatus.COMPLETED)
        resp = client.get(f"/jobs/{job_id}/download?format=xlsx", headers=_auth(token))
        assert resp.status_code == 400

    def test_completed_job_download(
        self, client: TestClient, sample_bytes: bytes, tmp_path: Path
    ) -> None:
        token = _login(client)
        create_resp = client.post(
            "/jobs",
            files={"file": ("test.docx", sample_bytes, "application/octet-stream")},
            headers=_auth(token),
        )
        job_id = create_resp.json()["job_id"]
        from uuid import UUID

        import docx

        out_dir = tmp_path / "output"
        out_dir.mkdir()
        docx_path = out_dir / "output.docx"
        doc = docx.Document()
        doc.add_paragraph("test")
        doc.save(str(docx_path))

        store = client.app.state.job_store  # type: ignore[union-attr]
        store.update_job(
            UUID(job_id),
            status=JobStatus.COMPLETED,
            output_docx_path=str(docx_path),
        )

        resp = client.get(f"/jobs/{job_id}/download?format=docx", headers=_auth(token))
        assert resp.status_code == 200
        assert "wordprocessingml" in resp.headers["content-type"]
        zf = zipfile.ZipFile(io.BytesIO(resp.content))
        assert "word/document.xml" in zf.namelist()


class TestJobListPagination:
    def _upload(
        self, client: TestClient, token: str, filename: str, content: bytes
    ) -> str:
        resp = client.post(
            "/jobs",
            files={"file": (filename, content, "application/octet-stream")},
            headers=_auth(token),
        )
        assert resp.status_code == 201
        return resp.json()["job_id"]  # type: ignore[no-any-return]

    def test_list_default_shape(self, client: TestClient, sample_bytes: bytes) -> None:
        token = _login(client)
        self._upload(client, token, "a.docx", sample_bytes)
        resp = client.get("/jobs", headers=_auth(token))
        assert resp.status_code == 200
        body = resp.json()
        assert "items" in body and "total" in body
        assert body["total"] == 1
        assert len(body["items"]) == 1

    def test_pagination_boundaries(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        token = _login(client)
        for i in range(5):
            self._upload(client, token, f"file{i}.docx", sample_bytes)

        resp = client.get("/jobs?limit=2&offset=0", headers=_auth(token))
        body = resp.json()
        assert body["total"] == 5
        assert len(body["items"]) == 2

        resp = client.get("/jobs?limit=2&offset=4", headers=_auth(token))
        body = resp.json()
        assert len(body["items"]) == 1  # only one left at the tail

        resp = client.get("/jobs?limit=2&offset=10", headers=_auth(token))
        body = resp.json()
        assert body["items"] == []
        assert body["total"] == 5

    def test_limit_capped_at_max(self, client: TestClient, sample_bytes: bytes) -> None:
        token = _login(client)
        self._upload(client, token, "a.docx", sample_bytes)
        resp = client.get("/jobs?limit=99999", headers=_auth(token))
        assert resp.status_code == 200  # doesn't error, just capped server-side

    def test_search_matches_filename_case_insensitive(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        token = _login(client)
        self._upload(client, token, "Informatika_Test.docx", sample_bytes)
        self._upload(client, token, "Riyaziyyat.docx", sample_bytes)

        resp = client.get("/jobs?search=informatika", headers=_auth(token))
        body = resp.json()
        assert body["total"] == 1
        assert body["items"][0]["input_filename"] == "Informatika_Test.docx"

    def test_search_no_match_returns_empty(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        token = _login(client)
        self._upload(client, token, "a.docx", sample_bytes)
        resp = client.get("/jobs?search=nonexistent", headers=_auth(token))
        body = resp.json()
        assert body["total"] == 0
        assert body["items"] == []

    def test_search_combined_with_pagination(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        token = _login(client)
        for i in range(3):
            self._upload(client, token, f"aligned_{i}.docx", sample_bytes)
        self._upload(client, token, "other.docx", sample_bytes)

        resp = client.get("/jobs?search=aligned&limit=2&offset=0", headers=_auth(token))
        body = resp.json()
        assert body["total"] == 3
        assert len(body["items"]) == 2

        resp = client.get("/jobs?search=aligned&limit=2&offset=2", headers=_auth(token))
        body = resp.json()
        assert len(body["items"]) == 1

    def test_ownership_scoping_respected_under_search(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        admin_token = _login(client)
        client.post(
            "/users",
            json={"email": "member@test.com", "password": "memberpass"},
            headers=_auth(admin_token),
        )
        member_token = _login(client, "member@test.com", "memberpass")

        self._upload(client, admin_token, "shared_name.docx", sample_bytes)
        self._upload(client, member_token, "shared_name.docx", sample_bytes)

        resp = client.get("/jobs?search=shared", headers=_auth(member_token))
        body = resp.json()
        assert body["total"] == 1  # member sees only their own match

        resp = client.get("/jobs?search=shared", headers=_auth(admin_token))
        body = resp.json()
        assert body["total"] == 2  # admin sees both


class TestJobDeletion:
    def test_owner_can_delete_own_job(
        self, client: TestClient, sample_bytes: bytes
    ) -> None:
        token = _login(client)
        create_resp = client.post(
            "/jobs",
            files={"file": ("test.docx", sample_bytes, "application/octet-stream")},
            headers=_auth(token),
        )
        job_id = create_resp.json()["job_id"]

        resp = client.delete(f"/jobs/{job_id}", headers=_auth(token))
        assert resp.status_code == 204

        resp = client.get(f"/jobs/{job_id}", headers=_auth(token))
        assert resp.status_code == 404

    def test_delete_removes_files_from_disk(
        self, client: TestClient, sample_bytes: bytes, tmp_path: Path
    ) -> None:
        token = _login(client)
        create_resp = client.post(
            "/jobs",
            files={"file": ("test.docx", sample_bytes, "application/octet-stream")},
            headers=_auth(token),
        )
        job_id = create_resp.json()["job_id"]

        from uuid import UUID

        import docx

        out_dir = tmp_path / "output"
        out_dir.mkdir()
        docx_path = out_dir / "output.docx"
        docx.Document().save(str(docx_path))

        store = client.app.state.job_store  # type: ignore[union-attr]
        store.update_job(
            UUID(job_id),
            status=JobStatus.COMPLETED,
            output_docx_path=str(docx_path),
        )
        assert docx_path.exists()

        resp = client.delete(f"/jobs/{job_id}", headers=_auth(token))
        assert resp.status_code == 204
        assert not docx_path.exists()

    def test_delete_nonexistent_job_404(self, client: TestClient) -> None:
        token = _login(client)
        resp = client.delete(
            "/jobs/00000000-0000-0000-0000-000000000000", headers=_auth(token)
        )
        assert resp.status_code == 404

    def test_unauthenticated_delete_401(self, client: TestClient) -> None:
        resp = client.delete("/jobs/00000000-0000-0000-0000-000000000000")
        assert resp.status_code == 401
