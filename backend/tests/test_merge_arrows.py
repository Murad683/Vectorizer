"""Arrows that merge into another arrow's line keep their own arrowhead.

Every textbook sample (samples/sp_open_exit.png, sum_loop.jpg,
quadratic_equation.png, flowchart_a3.png) draws a loop or a second branch
that joins the line into a symbol with an arrowhead AT the junction,
pointing at that line; the line itself carries on to the symbol with its
own arrowhead. Drawn in full, the joining line used to hide under the other
one and its arrowhead disappeared.
"""

import io
import zipfile

import docx
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app.core.layout_boxes import _cut_at_merges
from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.models.diagram import DiagramEdge
from tests.diagram_fixtures import run_fixture

_WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _edges(*pairs: tuple[str, str]) -> list[DiagramEdge]:
    return [DiagramEdge(source_id=s, target_id=t) for s, t in pairs]


def test_side_join_is_cut_at_the_junction() -> None:
    edges = _edges(("a", "t"), ("b", "t"))
    routed = {
        0: ([(0.0, 0.0), (0.0, 100.0)], False),  # straight trunk into t
        1: ([(50.0, 50.0), (0.0, 50.0), (0.0, 100.0)], True),  # joins from the right
    }
    _cut_at_merges(edges, routed)
    assert routed[0][0] == [(0.0, 0.0), (0.0, 100.0)]
    assert routed[1][0] == [(50.0, 50.0), (0.0, 50.0)]


def test_junction_in_the_middle_of_a_run_is_found() -> None:
    """The trunk turns onto the shared line further down; the joiner meets
    it in the middle of the trunk's straight run (A diagram: 8 -> Son)."""
    edges = _edges(("a", "t"), ("b", "t"))
    trunk = [(0.0, 0.0), (-40.0, 0.0), (-40.0, 200.0), (0.0, 200.0)]
    joiner = [(0.0, 100.0), (-40.0, 100.0), (-40.0, 200.0), (0.0, 200.0)]
    routed = {0: (trunk, False), 1: (joiner, False)}
    _cut_at_merges(edges, routed)
    assert routed[0][0] == trunk
    assert routed[1][0] == [(0.0, 100.0), (-40.0, 100.0)]


def test_lines_into_different_symbols_are_never_cut() -> None:
    edges = _edges(("a", "t"), ("b", "u"))
    routed = {
        0: ([(0.0, 0.0), (0.0, 100.0)], False),
        1: ([(50.0, 50.0), (0.0, 50.0), (0.0, 100.0)], False),
    }
    _cut_at_merges(edges, routed)
    assert len(routed[1][0]) == 3


def test_lines_turning_a_corner_together_are_not_a_junction() -> None:
    """Both run down the left and turn the same corner: neither meets a line
    that runs straight through, so neither is cut."""
    edges = _edges(("a", "t"), ("b", "t"))
    first = [(0.0, 0.0), (-40.0, 0.0), (-40.0, 200.0), (0.0, 200.0)]
    second = [(0.0, 0.0), (-40.0, 0.0), (-40.0, 200.0), (0.0, 200.0)]
    routed = {0: (list(first), False), 1: (list(second), False)}
    _cut_at_merges(edges, routed)
    assert routed[0][0] == first and routed[1][0] == second


def test_every_sample_junction_gets_its_own_arrow() -> None:
    expected = {
        "sp_open_exit": {("n5", "n2"), ("n6", "n7")},  # loop, P*3 -> open exit
        "sum_loop": {("n7", "n5")},  # loop onto the line into S = S + X(I)
        "quadratic_equation": {("n6", "n11"), ("n8", "n11")},  # into "Son"
    }
    for name, joiners in expected.items():
        _, lg = run_fixture(name)
        cut = set()
        for e in lg.edges:
            target = lg.node_by_id(e.target_id)
            assert e.edge_path is not None
            ex, ey = e.edge_path[-1]
            on_target = (
                target.x - 1 <= ex <= target.x + target.width + 1
                and target.y - 1 <= ey <= target.y + target.height + 1
            )
            if not on_target:
                cut.add((e.source_id, e.target_id))
        assert cut == joiners, name


def test_docx_and_pptx_draw_both_arrowheads() -> None:
    _, lg = run_fixture("sum_loop")
    document = docx.Document()
    add_diagram_to_docx(lg, document, max_width_emu=Inches(6), max_height_emu=Inches(7))
    buf = io.BytesIO()
    document.save(buf)
    with zipfile.ZipFile(buf) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    arrows = {}
    for wsp in root.iter(f"{_WPS}wsp"):
        name = wsp.find(f"{_WPS}cNvPr").get("name")
        if name.startswith("edge-"):
            arrows[name] = wsp.find(f".//{_A}tailEnd") is not None
    assert arrows["edge-n7-n5"] and arrows["edge-n4-n5"]

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(lg, slide, max_width_emu=Inches(9), max_height_emu=Inches(7))
    for name in ("edge-n7-n5", "edge-n4-n5"):
        shape = next(s for s in slide.shapes if s.name == name)
        assert shape._element.find(f".//{_A}tailEnd") is not None
