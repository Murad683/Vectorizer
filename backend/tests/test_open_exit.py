"""Open arrow ends (``ShapeType.EXIT``) and the topology rules around them.

Real case (samples/sp_open_exit.png): both ``P = P * k`` blocks run into one
arrow that stops in empty space. Forcing "every process has one outgoing
arrow" made the corrective retry invent ``P*2 -> P*3`` and a self-loop on
``P*3``; with an exit node the true drawing is structurally valid.
"""

import io
import zipfile

import docx
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app.core.diagram_quality import quality_issues
from app.core.layout import layout_graph
from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.core.vision.postprocess import topology_problems
from app.models.diagram import DiagramEdge, DiagramGraph, DiagramNode, ShapeType
from tests.diagram_fixtures import run_fixture, sp_open_exit

_WPS = "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def test_true_drawing_has_no_topology_problems() -> None:
    assert topology_problems(sp_open_exit()) == []


def test_missing_exit_is_still_reported() -> None:
    graph = sp_open_exit()
    graph.nodes = [n for n in graph.nodes if n.shape_type != ShapeType.EXIT]
    graph.edges = [e for e in graph.edges if e.target_id != "n7"]
    problems = topology_problems(graph)
    assert len(problems) == 2 and all("0 outgoing" in p for p in problems)


def test_invented_self_loop_is_reported() -> None:
    graph = sp_open_exit()
    graph.edges.append(DiagramEdge(source_id="n6", target_id="n6"))
    assert any("to itself" in p for p in topology_problems(graph))


def test_only_the_first_symbol_of_a_fragment_may_lack_an_incoming_arrow() -> None:
    graph = sp_open_exit()
    graph.edges = [e for e in graph.edges if e.target_id != "n4"]
    problems = topology_problems(graph)
    assert any("(n4) has no incoming arrow" in p for p in problems)
    assert not any("(n1)" in p for p in problems)


def test_exit_needs_an_incoming_arrow() -> None:
    graph = DiagramGraph(
        nodes=[
            DiagramNode(id="n1", shape_type=ShapeType.TERMINATOR, text="Başlanğıc"),
            DiagramNode(id="n2", shape_type=ShapeType.EXIT, text=""),
        ],
        edges=[],
    )
    assert any("open arrow end" in p for p in topology_problems(graph))


def test_exit_keeps_exact_mode_and_lands_on_the_drawn_arrow_tip() -> None:
    graph, lg = run_fixture("sp_open_exit")
    assert graph.boxes_detected  # the exit needs no detected symbol
    assert quality_issues(lg) == []
    end = lg.node_by_id("n7")
    assert end.width == 0 and end.height == 0
    into = [e for e in lg.edges if e.target_id == "n7"]
    assert len(into) == 2
    for e in into:
        assert e.edge_path is not None
        assert e.edge_path[-1] == [end.x, end.y]
    # Like the source: the open line runs down the right of every symbol.
    rightmost = max(n.x + n.width for n in lg.nodes if n.id != "n7")
    assert end.x > rightmost
    assert end.y > max(n.y + n.height for n in lg.nodes if n.id != "n7")


def test_docx_draws_no_symbol_and_free_lines_into_the_exit() -> None:
    _, lg = run_fixture("sp_open_exit")
    document = docx.Document()
    add_diagram_to_docx(lg, document, max_width_emu=Inches(6), max_height_emu=Inches(7))
    buf = io.BytesIO()
    document.save(buf)
    with zipfile.ZipFile(buf) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    names = [c.get("name") for c in root.iter(f"{_WPS}cNvPr")]
    assert "node-n7" not in names
    assert {"edge-n3-n7", "edge-n6-n7"} <= set(names)
    for cxn in root.iter(f"{_A}endCxn"):
        assert cxn.get("id") != "0"


def test_pptx_draws_free_lines_into_the_exit() -> None:
    _, lg = run_fixture("sp_open_exit")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(lg, slide, max_width_emu=Inches(9), max_height_emu=Inches(7))
    names = [s.name for s in slide.shapes]
    assert {"edge-n3-n7", "edge-n6-n7"} <= set(names)
    assert len([s for s in slide.shapes if s.shape_type == 1]) == 6  # symbols only


def test_graphviz_fallback_handles_an_exit() -> None:
    graph = sp_open_exit()
    for n in graph.nodes:
        n.box_2d = None  # no geometry at all -> Graphviz
    lg = layout_graph(graph)
    assert len(lg.edges) == len(graph.edges)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(lg, slide, max_width_emu=Inches(9), max_height_emu=Inches(7))
    assert "edge-n6-n7" in [s.name for s in slide.shapes]
