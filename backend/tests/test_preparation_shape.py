"""The preparation symbol (hexagon loop header, e.g. "i=2,N") and the two
real samples that came with it (offline)."""

import io
import zipfile

import docx
import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.core.diagram_detect import detect
from app.core.layout_boxes import _weld_corners
from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.core.vision.gemini_provider import _PROMPT
from app.core.vision.postprocess import topology_problems
from app.models.diagram import DiagramEdge, ShapeType
from tests.diagram_fixtures import SAMPLES, factorial_preparation, run_fixture

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.mark.parametrize("file", ["factorial_preparation.png", "xy_loop_photo.png"])
def test_every_symbol_is_detected(file: str) -> None:
    det = detect((SAMPLES / file).read_bytes())
    assert det is not None
    assert len(det.shapes) == 7


def test_loop_header_with_body_and_exit_is_valid() -> None:
    assert topology_problems(factorial_preparation()) == []


def test_loop_header_needs_one_or_two_exits() -> None:
    graph = factorial_preparation()
    graph.edges.append(DiagramEdge(source_id="n4", target_id="n7"))  # a third one
    assert any("loop header" in p for p in topology_problems(graph))
    graph = factorial_preparation()
    graph.edges = [e for e in graph.edges if e.source_id != "n4"]
    assert any("loop header" in p for p in topology_problems(graph))


def test_loop_enters_and_leaves_through_the_hexagon_points() -> None:
    """Like the source: the loop returns into the LEFT point, the exit
    leaves from the RIGHT point, the body hangs below."""
    _, lg = run_fixture("factorial_preparation")
    header = lg.node_by_id("n4")
    assert header.shape_type == ShapeType.PREPARATION
    cy = header.y + header.height / 2
    loop = next(e for e in lg.edges if e.source_id == "n5" and e.target_id == "n4")
    leave = next(e for e in lg.edges if e.source_id == "n4" and e.target_id == "n6")
    body = next(e for e in lg.edges if e.source_id == "n4" and e.target_id == "n5")
    assert loop.edge_path and leave.edge_path and body.edge_path
    assert loop.edge_path[-1] == pytest.approx([header.x, cy], abs=0.5)
    assert leave.edge_path[0] == pytest.approx([header.x + header.width, cy], abs=0.5)
    assert body.edge_path[0][1] == pytest.approx(header.y + header.height, abs=0.5)


def test_docx_and_pptx_draw_the_preparation_preset() -> None:
    _, lg = run_fixture("factorial_preparation")
    document = docx.Document()
    add_diagram_to_docx(lg, document, max_width_emu=Inches(6), max_height_emu=Inches(7))
    buf = io.BytesIO()
    document.save(buf)
    with zipfile.ZipFile(buf) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    presets = [g.get("prst") for g in root.iter(f"{_A}prstGeom")]
    assert presets.count("flowChartPreparation") == 1
    assert "diamond" not in presets

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(lg, slide, max_width_emu=Inches(9), max_height_emu=Inches(7))
    kinds = [s.auto_shape_type for s in slide.shapes if s.shape_type == 1]
    assert kinds.count(MSO_SHAPE.FLOWCHART_PREPARATION) == 1


def test_prompt_tells_hexagon_from_diamond() -> None:
    assert '"preparation" (HEXAGON' in _PROMPT
    assert "never call a hexagon a" in _PROMPT


def test_photo_loop_lands_on_the_connector_above_the_decision() -> None:
    """X, Y photo: the loop back from ``Y = Y + 5`` joins the connector into
    ``X <= Y`` with its own arrowhead (it used to enter the left corner)."""
    _, lg = run_fixture("xy_loop_photo")
    decision = lg.node_by_id("n3")
    loop = next(e for e in lg.edges if e.source_id == "n5" and e.target_id == "n3")
    assert loop.edge_path is not None
    end_x, end_y = loop.edge_path[-1]
    assert end_x == pytest.approx(decision.x + decision.width / 2, abs=0.5)
    assert end_y < decision.y - 1


def test_corner_of_the_author_lines_is_welded() -> None:
    """Measured ends differ by the line thickness; a spacing cut between
    them used to tear the loop corner apart."""
    vert = [(45.0, 181.0, 465.0)]
    horiz = [(179.6, 46.5, 175.0)]
    v, h = _weld_corners(vert, horiz)
    assert v[0][1] == pytest.approx(179.6)
    assert h[0][1] == pytest.approx(45.0)
