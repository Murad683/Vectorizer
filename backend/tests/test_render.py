"""Renderer tests: debug SVG integration + native PPTX shape generation."""

import html
import xml.etree.ElementTree as ET
from pathlib import Path
from unittest.mock import MagicMock, patch

from pptx import Presentation
from pptx.shapes.connector import Connector
from pptx.util import Inches

from app.core.classifier import classify_elements
from app.core.ingestion.docx_extractor import DocxExtractor
from app.core.layout import layout_graph
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.core.render.svg_debug import render_svg
from app.models.diagram import DiagramEdge, DiagramGraph, DiagramNode, ShapeType
from app.models.document import ElementType, VisionResult


class TestSvgDebugRenderer:
    def test_full_pipeline_produces_valid_svg(
        self,
        sample_docx: Path,
        flowchart_png_bytes: bytes,
        flowchart_graph: DiagramGraph,
        tmp_path: Path,
    ) -> None:
        """Extract → classify → (mocked) vision → layout → SVG, on the
        real sample. Saved to a file for manual visual inspection."""
        elements = DocxExtractor().extract(sample_docx)
        classified = classify_elements(elements)
        diagram = next(
            c
            for c in classified
            if c.element_type == ElementType.DIAGRAM
            and c.element.raw_content == flowchart_png_bytes
        )
        mock_provider = MagicMock()
        mock_provider.analyze_diagram.return_value = VisionResult(
            diagram_graph=flowchart_graph
        )
        with patch("app.core.vision.get_vision_provider", return_value=mock_provider):
            from app.core.vision import get_vision_provider

            image = diagram.element.raw_content
            assert isinstance(image, bytes)
            result = get_vision_provider().analyze_diagram(image)

        assert result.diagram_graph is not None
        layouted = layout_graph(result.diagram_graph)
        svg = render_svg(layouted)

        out_file = tmp_path / "flowchart_debug.svg"
        out_file.write_text(svg, encoding="utf-8")

        root = ET.fromstring(svg)  # valid XML or this raises
        assert root.tag.endswith("svg")
        for node in flowchart_graph.nodes:
            assert html.escape(node.text) in svg
        # One drawn symbol per node (rect/polygon), plus the background rect.
        ns = "{http://www.w3.org/2000/svg}"
        symbols = root.findall(f".//{ns}rect") + root.findall(f".//{ns}polygon")
        assert len(symbols) - 1 == len(flowchart_graph.nodes)

    def test_back_edge_and_labels_present(self, flowchart_graph: DiagramGraph) -> None:
        svg = render_svg(layout_graph(flowchart_graph))
        assert "Hə" in svg
        assert "Yox" in svg
        assert "&#8635;" in svg  # loop icon marks the routed back-edge


class TestPptxShapes:
    def _build_pptx(self, flowchart_graph: DiagramGraph, path: Path) -> Path:
        layouted = layout_graph(flowchart_graph)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        add_diagram_to_slide(
            layouted,
            slide,
            offset_x_emu=Inches(0.5),
            offset_y_emu=Inches(0.2),
            max_width_emu=Inches(9),
            max_height_emu=Inches(7),
        )
        prs.save(str(path))
        return path

    def test_shapes_are_native_and_text_editable(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "diagram.pptx")
        prs = Presentation(str(path))  # reopen from disk
        slide = prs.slides[0]
        autoshapes = [
            s for s in slide.shapes if s.has_text_frame and not s.is_placeholder
        ]
        texts = {s.text_frame.text for s in autoshapes if s.text_frame.text}
        for node in flowchart_graph.nodes:
            assert node.text in texts  # Azerbaijani text survived round-trip

    def test_connectors_are_attached_to_shapes(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "diagram.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        connectors = [s for s in slide.shapes if isinstance(s, Connector)]
        assert len(connectors) == len(flowchart_graph.edges)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for connector in connectors:
            assert connector._element.find(".//a:stCxn", ns) is not None
            assert connector._element.find(".//a:endCxn", ns) is not None

    def test_parallelogram_uses_custom_geometry_bbox_centered(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "diagram.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        io_node_ids = [
            n.id for n in flowchart_graph.nodes if n.shape_type.value == "input_output"
        ]
        assert io_node_ids
        shapes = [
            s
            for s in slide.shapes
            if s.has_text_frame
            and s.text_frame.text
            in {n.text for n in flowchart_graph.nodes if n.id in io_node_ids}
        ]
        assert len(shapes) == len(io_node_ids)
        for shape in shapes:
            prst_geom = shape._element.find(".//a:prstGeom", ns)
            assert prst_geom is None
            cust_geom = shape._element.find(".//a:custGeom", ns)
            assert cust_geom is not None
            cxns = cust_geom.findall("a:cxnLst/a:cxn/a:pos", ns)
            assert len(cxns) == 4
            rect = cust_geom.find("a:rect", ns)
            width = int(rect.get("r"))
            assert int(cxns[0].get("x")) == width // 2  # top
            assert int(cxns[2].get("x")) == width // 2  # bottom

    def test_back_edge_is_glued_and_routed_from_graphviz_path(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        """Phase 24: back-edges are no longer forced onto a fixed
        right-side "rail" — they route through the same
        choose_connector_from_path() pipeline as forward edges, using
        Graphviz's own collision-free path. Connection sites and preset
        now depend on actual geometry rather than being hardcoded to
        idx=3 (right) on both ends."""
        layouted = layout_graph(flowchart_graph)
        back_edges = [e for e in layouted.edges if e.is_back_edge]
        assert len(back_edges) == 1
        path = self._build_pptx(flowchart_graph, tmp_path / "diagram.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        connectors = [s for s in slide.shapes if isinstance(s, Connector)]
        # One connector per edge, all glued.
        assert len(connectors) == len(flowchart_graph.edges)
        for conn in connectors:
            assert conn._element.find(".//a:stCxn", ns) is not None
            assert conn._element.find(".//a:endCxn", ns) is not None

    def test_edge_labels_become_text_boxes(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "diagram.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        all_text = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        assert "Hə" in all_text
        assert "Yox" in all_text

    def test_edge_label_does_not_overlap_source_decision(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        decision_text = next(n.text for n in flowchart_graph.nodes if n.id == "n4")
        path = self._build_pptx(flowchart_graph, tmp_path / "diagram.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        decision_shape = next(
            s
            for s in slide.shapes
            if s.has_text_frame and s.text_frame.text == decision_text
        )
        d_left = decision_shape.left
        d_right = decision_shape.left + decision_shape.width
        d_top = decision_shape.top
        d_bottom = decision_shape.top + decision_shape.height
        for label_text in ("Hə", "Yox"):
            label_shape = next(
                s
                for s in slide.shapes
                if s.has_text_frame and s.text_frame.text == label_text
            )
            l_left = label_shape.left
            l_right = label_shape.left + label_shape.width
            l_top = label_shape.top
            l_bottom = label_shape.top + label_shape.height
            no_overlap = (
                l_right <= d_left
                or l_left >= d_right
                or l_bottom <= d_top
                or l_top >= d_bottom
            )
            assert no_overlap, (
                f"Label '{label_text}' overlaps decision shape"
            )


class TestBackEdgePptx:
    """All back-edges must be glued native connectors — no freeform detours.

    Phase 24: back-edges route through choose_connector_from_path() just
    like forward edges, using Graphviz's own collision-free path — so
    the resulting preset/connection-sites depend on actual geometry
    rather than being hardcoded to a fixed "rail" (bentConnector3,
    idx=3/3 on both ends)."""

    def _build_pptx(self, graph: DiagramGraph, path: Path) -> Path:
        layouted = layout_graph(graph)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_diagram_to_slide(
            layouted,
            slide,
            offset_x_emu=Inches(0.5),
            offset_y_emu=Inches(0.2),
            max_width_emu=Inches(9),
            max_height_emu=Inches(7),
        )
        prs.save(str(path))
        return path

    def test_all_back_edges_glued_no_freeform(
        self, collision_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        layouted = layout_graph(collision_graph)
        back_edge_count = sum(e.is_back_edge for e in layouted.edges)
        assert back_edge_count == 2
        path = self._build_pptx(collision_graph, tmp_path / "detour.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        freeforms = [
            s
            for s in slide.shapes
            if not isinstance(s, Connector)
            and s._element.find(".//a:custGeom", ns) is not None
            and s._element.find(".//a:custGeom//a:path", ns) is not None
            and s._element.find(".//a:custGeom//a:path", ns).get("fill")
            == "none"
        ]
        assert len(freeforms) == 0
        connectors = [s for s in slide.shapes if isinstance(s, Connector)]
        assert len(connectors) == len(collision_graph.edges)
        for conn in connectors:
            assert conn._element.find(".//a:stCxn", ns) is not None
            assert conn._element.find(".//a:endCxn", ns) is not None


class TestStructuralVerificationPptx:
    """Verify that the rendered PPTX's graph structure exactly matches
    the input graph."""

    def _build_pptx(self, graph: DiagramGraph, path: Path) -> Path:
        layouted = layout_graph(graph)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_diagram_to_slide(
            layouted,
            slide,
            offset_x_emu=Inches(0.5),
            offset_y_emu=Inches(0.2),
            max_width_emu=Inches(9),
            max_height_emu=Inches(7),
        )
        prs.save(str(path))
        return path

    def test_all_node_texts_present(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "verify.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        rendered_texts = {
            s.text_frame.text
            for s in slide.shapes
            if s.has_text_frame and s.text_frame.text
        }
        for node in flowchart_graph.nodes:
            assert node.text in rendered_texts, f"Node '{node.text}' missing"

    def test_correct_number_of_connectors(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "verify.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        connectors = [s for s in slide.shapes if isinstance(s, Connector)]
        assert len(connectors) == len(flowchart_graph.edges)

    def test_all_connectors_glued_at_both_ends(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "verify.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        connectors = [s for s in slide.shapes if isinstance(s, Connector)]
        for conn in connectors:
            assert conn._element.find(".//a:stCxn", ns) is not None
            assert conn._element.find(".//a:endCxn", ns) is not None

    def test_all_edge_labels_present(
        self, flowchart_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(flowchart_graph, tmp_path / "verify.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        rendered_texts = {
            s.text_frame.text
            for s in slide.shapes
            if s.has_text_frame and s.text_frame.text
        }
        for edge in flowchart_graph.edges:
            if edge.label:
                assert edge.label in rendered_texts

    def test_collision_graph_structure_complete(
        self, collision_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        path = self._build_pptx(collision_graph, tmp_path / "verify_coll.pptx")
        prs = Presentation(str(path))
        slide = prs.slides[0]
        rendered_texts = {
            s.text_frame.text
            for s in slide.shapes
            if s.has_text_frame and s.text_frame.text
        }
        for node in collision_graph.nodes:
            assert node.text in rendered_texts
        connectors = [s for s in slide.shapes if isinstance(s, Connector)]
        assert len(connectors) == len(collision_graph.edges)
        for edge in collision_graph.edges:
            if edge.label:
                assert edge.label in rendered_texts


class TestLabelPositioningPptx:
    """Edge labels must not overlap each other and must not overlap
    the source decision node. Graphviz handles placement."""

    def test_two_sibling_labels_distinct_in_pptx_output(
        self, collision_graph: DiagramGraph, tmp_path: Path
    ) -> None:
        layouted = layout_graph(collision_graph)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_diagram_to_slide(
            layouted,
            slide,
            offset_x_emu=Inches(0.5),
            offset_y_emu=Inches(0.2),
            max_width_emu=Inches(9),
            max_height_emu=Inches(7),
        )
        prs.save(str(tmp_path / "labels.pptx"))
        prs2 = Presentation(str(tmp_path / "labels.pptx"))
        slide2 = prs2.slides[0]
        label_positions: dict[str, tuple[int, int]] = {}
        for s in slide2.shapes:
            if s.has_text_frame and s.text_frame.text in ("Hə", "Yox"):
                label_positions[s.text_frame.text] = (s.left, s.top)
        assert "Hə" in label_positions and "Yox" in label_positions
        assert label_positions["Hə"] != label_positions["Yox"], (
            f"Labels at identical position: {label_positions['Hə']}"
        )
