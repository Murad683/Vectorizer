"""Sub/superscript markup, the LaTeX subset and Word equations (offline)."""

import io
import zipfile

import docx
import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app.core.render.docx_shapes import add_diagram_to_docx
from app.core.render.pptx_shapes import add_diagram_to_slide
from app.core.rich_text import (
    M_NS,
    LatexError,
    display_lines,
    latex_runs,
    latex_to_omml,
    latex_to_text,
    markup_runs,
    plain,
)
from tests.diagram_fixtures import SAMPLES, run_fixture

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS = {"m": M_NS, "w": _W}
_QUAD_ROOTS = r"x_1 = \frac{-b + \sqrt{D}}{2a}, \quad x_2 = \frac{-b - \sqrt{D}}{2a}"


def test_markup_runs() -> None:
    assert markup_runs("X_{1}, X_{20} ədədləri") == [
        ("X", None),
        ("1", "sub"),
        (", X", None),
        ("20", "sub"),
        (" ədədləri", None),
    ]
    assert markup_runs("b^{2} - 4ac") == [("b", None), ("2", "sup"), (" - 4ac", None)]
    assert markup_runs("x_1") == [("x", None), ("1", "sub")]
    assert plain("max_val = X_{1}") == "max_val = X1"  # a bare _word stays text


def test_latex_linear_forms() -> None:
    assert latex_to_text(r"x = \frac{-b}{2a}") == "x = (-b)/(2a)"
    assert latex_to_text(r"\sqrt{D} \le 5") == "√D ≤ 5"
    assert latex_runs(r"S = x_1 + x_{20}") == [
        ("S = x", None),
        ("1", "sub"),
        (" + x", None),
        ("20", "sub"),
    ]


@pytest.mark.parametrize("bad", [r"\frac{1}", r"\unknowncmd{x}", r"{x", r"x}"])
def test_unsupported_latex_raises(bad: str) -> None:
    with pytest.raises(LatexError):
        latex_to_text(bad)


def test_omml_structure() -> None:
    xml = latex_to_omml(_QUAD_ROOTS, 24)
    root = etree.fromstring(f'<w:p xmlns:w="{_W}" xmlns:m="{M_NS}">{xml}</w:p>')
    assert len(root.findall(".//m:f", _NS)) == 2  # two stacked fractions
    assert len(root.findall(".//m:rad", _NS)) == 2  # two square roots
    assert len(root.findall(".//m:sSub", _NS)) == 2  # x1, x2
    texts = "".join(t.text or "" for t in root.iter(f"{{{M_NS}}}t"))
    assert "−b" in texts  # typeset with a real minus sign


def test_fraction_counts_as_two_lines_and_is_narrower() -> None:
    chars, lines = display_lines("x = -b / 2a", r"x = \frac{-b}{2a}")
    assert lines == 2 and chars < len("x = (-b)/(2a)") * 1.5
    assert display_lines("X_{1}, X_{2}\ndaxil") == (6, 2)


def _docx_root(name: str) -> etree._Element:
    _, lg = run_fixture(name)
    document = docx.Document()
    add_diagram_to_docx(lg, document, max_width_emu=Inches(6), max_height_emu=Inches(7))
    buf = io.BytesIO()
    document.save(buf)
    with zipfile.ZipFile(buf) as z:
        return etree.fromstring(z.read("word/document.xml"))


def test_docx_formulas_are_word_equations() -> None:
    root = _docx_root("quadratic_equation")
    maths = root.findall(".//m:oMath", _NS)
    assert len(maths) == 2
    assert len(root.findall(".//m:f", _NS)) == 3  # -b/2a and the two roots
    subs = root.findall(".//w:vertAlign[@w:val='subscript']", _NS)
    assert len(subs) == 2  # "x1 və x2-nin ..." print block


def test_docx_caption_keeps_subscripts() -> None:
    from app.core.assembly.docx_assembler import assemble_docx
    from app.models.document import (
        ClassifiedElement,
        ElementKind,
        ElementType,
        ExtractedElement,
    )
    from scripts.replay_diagram import _FixedProvider

    graph, _ = run_fixture("sum_loop")
    element = ClassifiedElement(
        element=ExtractedElement(
            source_page=1,
            element_kind=ElementKind.IMAGE,
            raw_content=(SAMPLES / "sum_loop.jpg").read_bytes(),
        ),
        element_type=ElementType.DIAGRAM,
        rule="test",
    )
    document = assemble_docx([element], vision_provider=_FixedProvider(graph))
    caption = next(p for p in document.paragraphs if p.text.startswith("Sxem 3."))
    lowered = [r.text for r in caption.runs if r.font.subscript]
    assert lowered == ["1", "2", "20"]
    assert caption.text.startswith("Sxem 3. S = x1 + x2")


def test_pptx_formula_has_real_subscripts() -> None:
    _, lg = run_fixture("quadratic_equation")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_diagram_to_slide(lg, slide, max_width_emu=Inches(9), max_height_emu=Inches(7))
    roots = next(
        s for s in slide.shapes if s.has_text_frame and "√D" in s.text_frame.text
    )
    lowered = [
        r.text
        for p in roots.text_frame.paragraphs
        for r in p.runs
        if r.font._rPr.get("baseline") == "-25000"
    ]
    assert lowered == ["1", "2"]
